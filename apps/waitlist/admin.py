from django.contrib import admin
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import WaitlistEntry

COLUMNS = [
    ("full_name", "الاسم / Name", 28),
    ("email", "البريد الإلكتروني / Email", 32),
    ("locale", "اللغة / Locale", 12),
    ("source", "المصدر / Source", 16),
    ("created_at", "تاريخ التسجيل / Joined At", 20),
]


@admin.action(description="تصدير المحدَّد كملف Excel / Export selected as Excel")
def export_waitlist_as_excel(modeladmin, request, queryset):
    """Admin action: download the selected waitlist rows as a formatted
    .xlsx (bold header, frozen top row, filter dropdowns, sensible column
    widths) -- this is the "professional Excel file" deliverable the
    marketing site's waitlist ultimately needs, layered on top of the real
    database rather than replacing it (writing straight to a shared .xlsx
    on every signup would corrupt under concurrent requests)."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Waitlist"

    header_fill = PatternFill(start_color="14172B", end_color="14172B", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    for col_index, (_, label, width) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=col_index, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(col_index)].width = width
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 26

    for row_index, entry in enumerate(queryset.order_by("-created_at"), start=2):
        sheet.cell(row=row_index, column=1, value=entry.full_name or "—")
        sheet.cell(row=row_index, column=2, value=entry.email)
        sheet.cell(row=row_index, column=3, value=entry.get_locale_display())
        sheet.cell(row=row_index, column=4, value=entry.get_source_display())
        joined_cell = sheet.cell(row=row_index, column=5, value=timezone.localtime(entry.created_at).strftime("%Y-%m-%d %H:%M"))
        joined_cell.alignment = Alignment(horizontal="center")

    last_row = max(sheet.max_row, 1)
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last_row}"

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    filename = f"baraq-waitlist-{timezone.localdate().isoformat()}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    workbook.save(response)
    return response


@admin.register(WaitlistEntry)
class WaitlistEntryAdmin(admin.ModelAdmin):
    list_display = ("full_name", "email", "locale", "source", "created_at")
    list_filter = ("locale", "source")
    search_fields = ("email", "full_name")
    ordering = ("-created_at",)
    readonly_fields = ("created_at", "updated_at")
    actions = [export_waitlist_as_excel]
